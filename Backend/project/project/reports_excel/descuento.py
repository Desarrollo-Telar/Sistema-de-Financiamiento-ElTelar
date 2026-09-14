from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from django.views.generic import TemplateView
from django.http import HttpResponse
from django.db.models import Q
from apps.financings.models import Descuento  # Ajusta el import según tu estructura

class ReporteDescuentosExcelView(TemplateView):
    """
    Genera un reporte en Excel de los Descuentos aplicados, 
    incluyendo datos actuales del descuento y la fotograma histórica guardada en data_cuota.
    """
    def get_queryset(self):
        try:
            request = self.request
            sucursal_id = request.session.get('sucursal_id')
            
            # Filtros por URL (opcionales)
            mes = request.GET.get('mes')
            anio = request.GET.get('anio')

            filters = Q()
            if sucursal_id:
                filters &= Q(sucursal_id=sucursal_id)
            if anio:
                filters &= Q(fecha_descuento__year=anio)
            if mes:
                filters &= Q(fecha_descuento__month=mes)

            # Optimización de consultas relacionando FKs principales
            return Descuento.objects.filter(filters).select_related(
                'credit', 'cuota', 'sucursal', 'usuario_descuento'
            ).order_by('-fecha_descuento')

        except Exception as e:
            print(f"Error en queryset de descuentos: {e}")
            return Descuento.objects.none()

    def get(self, request, *args, **kwargs):
        # 1. Crear el libro y la hoja de cálculo
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Reporte de Descuentos"

        # 2. Definición de encabezados
        headers = [
            # --- Datos del Descuento ---
            "#", 
            "Ref. Descuento (UUID)", 
            "Fecha Descuento", 
            "Crédito", 
            "Tipo Descuento", 
            "Motivo Descuento", 
            "Interés por Cobrar", 
            "Mora por Cobrar", 
            "Saldo Capital por Cobrar", 
            "Recalcular Cuota", 
            "Usuario Aplicó", 
            "Sucursal",
            "Activo",

            # --- Datos Históricos de la Cuota (Extraídos de data_cuota JSON) ---
            "Cuota (No. Mes)", 
            "Fecha Inicio Cuota", 
            "Fecha Vencimiento Cuota", 
         
            "Capital (Histórico)", 
            "Interés (Histórico)", 
            "Mora (Histórico)", 
            
            "Saldo Pendiente (Histórico)", 
           
          
        ]

        # Estilo opcional para los encabezados
        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
        
        sheet.append(headers)

        # Aplicar estilo a la fila de encabezado
        for cell in sheet[1]:
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")

        # 3. Iterar los datos y construir las filas
        descuentos = self.get_queryset()

        for idx, desc in enumerate(descuentos, start=1):
            # Extraer dict seguro de data_cuota (por si viniera None)
            data_c = desc.data_cuota or {}

            # Obtener datos formateados
            fecha_desc = desc.fecha_descuento.strftime("%Y-%m-%d %H:%M:%S") if desc.fecha_descuento else "---"
            
            # Fechas dentro del JSON (data_cuota)
            start_date_c = data_c.get('start_date', '---')
            due_date_c = data_c.get('due_date', '---')
            
            # Cortar la cadena ISO si viene completa (ej: "2026-09-14T00:00:00")
            if isinstance(start_date_c, str) and "T" in start_date_c:
                start_date_c = start_date_c.split("T")[0]
            if isinstance(due_date_c, str) and "T" in due_date_c:
                due_date_c = due_date_c.split("T")[0]

            sheet.append([
                idx,
                str(desc.numero_referencia),
                fecha_desc,
                str(desc.credit_id if hasattr(desc, 'credit_id') else desc.credit.id),
                str(desc.tipo_descuento),
                str(desc.motivo_descuento),
                float(desc.interes_por_cobrar or 0),
                float(desc.mora_por_cobrar or 0),
                float(desc.saldo_capital_por_cobrar or 0),
                "SÍ" if desc.recalcular_cuota else "NO",
                str(desc.usuario_descuento.get_full_name() or desc.usuario_descuento.username),
                str(desc.sucursal.name if hasattr(desc.sucursal, 'name') else desc.sucursal),
                "SÍ" if desc.activo else "NO",

                # --- Datos parseados del JSON `data_cuota` ---
                data_c.get('mes', '---'),
                start_date_c,
                due_date_c,
             
                data_c.get('principal', 0.0),
                data_c.get('interest', 0.0),
                data_c.get('mora', 0.0),
             
                data_c.get('saldo_pendiente', 0.0),
                
            ])

        # 4. Ajustar ancho de columnas automáticamente
        for col in sheet.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = col[0].column_letter
            sheet.column_dimensions[col_letter].width = max(max_len + 3, 12)

        # 5. Configurar respuesta HTTP
        response = HttpResponse(
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        mes = request.GET.get('mes') or "todos"
        anio = request.GET.get('anio') or "todos"
        response['Content-Disposition'] = f'attachment; filename="reporte_descuentos_{mes}_{anio}.xlsx"'

        workbook.save(response)
        return response